# research/tools/threshold_extract.py -- reads the owner's spreadsheets in threshold_strategy/ and writes
# what ALTAIR reads from them as plain CSV, under threshold_strategy/extracted/:
#   2D-H-L_trade_log.csv     the workbook's own trade log (sheet "Trade Log"), which altair_threshold
#                            checks its prev-2-day breakout against, trade by trade;
#   2D-H-L_days.csv          the workbook's sessions and closes (sheet "Equity Curve");
#   tradingview_backtests.csv one row per TradingView export in BACKTEST/, its own figures, unchanged
#                            (Performance, Trades analysis, Risk performance ratios, Properties); a *_frac or
#                            percent_profitable column is a fraction as TradingView writes it (3.51 = 351 %).
# Nothing is recomputed here. Run from the repo root after the spreadsheets change:
#   python3 research/tools/threshold_extract.py .
# Needs openpyxl (pip install openpyxl).
import csv, os, sys

import openpyxl

root = sys.argv[1] if len(sys.argv) > 1 else "."
src = os.path.join(root, "threshold_strategy")
out = os.path.join(src, "extracted")
os.makedirs(out, exist_ok=True)


def day(v):
    return v.date().isoformat() if hasattr(v, "date") else str(v)


# ---- the 2-day breakout workbook's trade log ----------------------------------------------------
wb = openpyxl.load_workbook(os.path.join(src, "market bnf sheets", "2D-H-L.xlsx"), read_only=True, data_only=True)
rows = list(wb["Trade Log"].iter_rows(values_only=True))
head = list(rows[0])
col = {name: head.index(name) for name in head if name is not None}
with open(os.path.join(out, "2D-H-L_trade_log.csv"), "w", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(["n", "side", "entry_date", "exit_date", "entry", "stop", "exit", "exit_reason", "days", "mfe_pts", "mae_pts",
                "gross_pts", "cost_pts", "net_pts", "flip"])
    n = 0
    for r in rows[1:]:
        if r[0] is None:
            continue
        g = lambda k: r[col[k]]
        w.writerow([g("#"), g("Direction"), day(g("Entry Date")), day(g("Exit Date")), g("Entry Price"), g("Stop Level"),
                    g("Exit Price"), g("Exit Reason"), g("Days Held"), g("MFE (pts)"), g("MAE (pts)"), g("Gross P&L (pts)"),
                    g("Cost (pts)"), g("Net P&L (pts)"), "Yes" if g("Flip?") == "Yes" else ""])
        n += 1
print(f"2D-H-L_trade_log.csv: {n} trades")
# Its sessions (sheet "Equity Curve"): the check runs once more on exactly these days.
with open(os.path.join(out, "2D-H-L_days.csv"), "w", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(["date", "close", "position"])
    n = 0
    for r in list(wb["Equity Curve"].iter_rows(values_only=True))[1:]:
        if r[0] is None:
            continue
        w.writerow([day(r[0]), r[2], r[1]])
        n += 1
print(f"2D-H-L_days.csv: {n} sessions")

# ---- TradingView exports -------------------------------------------------------------------------
fields = [
    ("symbol", "Properties", "Symbol", 1), ("timeframe", "Properties", "Timeframe", 1),
    ("trading_range", "Properties", "Trading range", 1), ("order_size", "Properties", "Order size", 1),
    ("point_value", "Properties", "Point value", 1), ("commission", "Properties", "Commission", 1),
    ("slippage", "Properties", "Slippage", 1), ("initial_capital", "Properties", "Initial capital", 1),
    ("total_trades", "Trades analysis", "Total trades", 1), ("percent_profitable", "Trades analysis", "Percent profitable", 2),
    ("net_profit_inr", "Performance", "Net profit", 1), ("net_profit_frac", "Performance", "Net profit", 2),
    ("gross_profit_inr", "Performance", "Gross profit", 1), ("gross_loss_inr", "Performance", "Gross loss", 1),
    ("commission_paid_inr", "Performance", "Commission paid", 1),
    ("max_drawdown_inr", "Performance", "Max equity drawdown", 1), ("max_drawdown_frac", "Performance", "Max equity drawdown", 2),
    ("buy_hold_inr", "Performance", "Buy & hold return", 1), ("profit_factor", "Risk performance ratios", "Profit factor", 1),
    ("sharpe", "Risk performance ratios", "Sharpe ratio", 1), ("sortino", "Risk performance ratios", "Sortino ratio", 1),
    ("avg_pnl_inr", "Trades analysis", "Avg P&L", 1), ("avg_win_inr", "Trades analysis", "Avg winning trade", 1),
    ("avg_loss_inr", "Trades analysis", "Avg losing trade", 1),
    ("largest_loss_inr", "Trades analysis", "Largest losing trade", 1),
    ("avg_bars", "Trades analysis", "Avg # bars in trades", 1),
]
bt = os.path.join(src, "BACKTEST")
files = sorted(f for f in os.listdir(bt) if f.endswith(".xlsx") and not f.startswith("._"))
with open(os.path.join(out, "tradingview_backtests.csv"), "w", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    # Counted from TradingView's own "List of trades", not recomputed: how many
    # trades enter and leave inside one bar (the same bar time). Inside a bar
    # TradingView does not know whether the high or the low came first (bar
    # magnifier off), so such trades rest on its guess.
    w.writerow([k for k, *_ in fields] + ["trades_listed", "same_bar_trades", "file"])
    for name in files:
        book = openpyxl.load_workbook(os.path.join(bt, name), read_only=True, data_only=True)
        sheets = {}
        for s in {sheet for _, sheet, _, _ in fields}:
            sheets[s] = {r[0]: r for r in book[s].iter_rows(values_only=True) if r and r[0] is not None}
        row = []
        for _, sheet, label, i in fields:
            r = sheets[sheet].get(label)
            v = r[i] if r is not None and i < len(r) else None
            row.append("" if v is None else v)
        legs = [r for r in book["List of trades"].iter_rows(values_only=True)][1:]
        times = {}
        for r in legs:
            if r and r[0] is not None and r[3] is not None:
                times.setdefault(r[0], []).append(r[3])
        same = sum(1 for t in times.values() if len(t) == 2 and t[0] == t[1])
        w.writerow(row + [len(times), same, name])
print(f"tradingview_backtests.csv: {len(files)} exports")
